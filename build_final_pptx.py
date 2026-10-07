import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_final_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palet Warna Cerah Bersih & Kontras (Aman untuk Canva)
    C_BG            = RGBColor(250, 248, 245)   # #FAF8F5 Warm Paper
    C_WHITE         = RGBColor(255, 255, 255)   # Pure White Card
    C_DARK          = RGBColor(15, 23, 42)      # #0F172A Dark Slate Text & Border
    C_MUTED         = RGBColor(71, 85, 105)     # #475569 Muted Text

    # Pastel Backgrounds
    C_PINK          = RGBColor(255, 228, 230)   # #FFE4E6 Rose Pastel
    C_BLUE          = RGBColor(224, 242, 254)   # #E0F2FE Sky Pastel
    C_YELLOW        = RGBColor(254, 240, 138)   # #FEF08A Yellow Pastel
    C_GREEN         = RGBColor(220, 252, 231)   # #DCFCE7 Mint Pastel
    C_PURPLE        = RGBColor(243, 232, 255)   # #F3E8FF Lilac Pastel
    C_ORANGE        = RGBColor(254, 215, 170)   # #FED7AA Peach Pastel

    FONT_FAMILY     = "Arial"

    def add_base_slide(slide, url_path, badge_text):
        # 1. Background Canvas
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()

        # 2. Main Window Box (White Container)
        win = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.4), Inches(12.133), Inches(6.7))
        win.fill.solid()
        win.fill.fore_color.rgb = C_WHITE
        win.line.color.rgb = C_DARK
        win.line.width = Pt(2)

        # 3. Top Window Bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.4), Inches(12.133), Inches(0.6))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
        bar.line.color.rgb = C_DARK
        bar.line.width = Pt(1.5)

        # 4. Three Dots (Traffic Lights)
        colors = [RGBColor(255, 95, 86), RGBColor(255, 189, 46), RGBColor(39, 201, 63)]
        for i, c in enumerate(colors):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9 + i * 0.28), Inches(0.58), Inches(0.18), Inches(0.18))
            dot.fill.solid()
            dot.fill.fore_color.rgb = c
            dot.line.color.rgb = C_DARK
            dot.line.width = Pt(1)

        # 5. URL Address Pill
        url_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(0.49), Inches(7.8), Inches(0.4))
        url_pill.fill.solid()
        url_pill.fill.fore_color.rgb = C_WHITE
        url_pill.line.color.rgb = C_DARK
        url_pill.line.width = Pt(1.5)
        tf_u = url_pill.text_frame
        tf_u.word_wrap = False
        p_u = tf_u.paragraphs[0]
        p_u.text = f"https://{url_path}"
        p_u.font.name = FONT_FAMILY
        p_u.font.size = Pt(10)
        p_u.font.bold = True
        p_u.font.color.rgb = C_DARK
        p_u.alignment = PP_ALIGN.CENTER

        # 6. Status Badge Pill Right
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.49), Inches(2.3), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_YELLOW
        badge.line.color.rgb = C_DARK
        badge.line.width = Pt(1.5)
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = C_DARK
        p_b.alignment = PP_ALIGN.CENTER

    def add_header(slide, tag, title, subtitle=''):
        tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(11.5), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"✦ {tag.upper()}"
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(147, 51, 234)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = C_DARK

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = FONT_FAMILY
            p2.font.size = Pt(11)
            p2.font.color.rgb = C_MUTED

    def add_card(slide, left, top, width, height, bg_color):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = C_DARK
        card.line.width = Pt(2)
        return card

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_base_slide(s1, "sisapras.jtik.unm.ac.id/portal", "KELOMPOK 2 (HCI)")

    tb1 = s1.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(6.5), Inches(5.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "PROPOSAL SEMINAR UTS • DESAIN ANTARMUKA PENGGUNA"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(219, 39, 119)

    p = tf1.add_paragraph()
    p.text = "SI-SAPRAS JTIK"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    p = tf1.add_paragraph()
    p.text = "Rancang Bangun Antarmuka Pengguna Sistem Informasi Peminjaman Sarana Prasarana di JTIK Berbasis Human-Centered Design (HCD)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = C_MUTED

    p = tf1.add_paragraph()
    p.text = "\nTim Pengembang (Kelompok 2 HCI):\n• Asyifa Nur Cahya Kamila (250210500014)\n• Sitti Aisyah Nur Azizah (250210500030)\n• Warnisyah Armin (250210500034)\n• Rezky Awalya (250210500043)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK

    p = tf1.add_paragraph()
    p.text = "\nDosen Pengampu: Ayu Lestari, S.Pd., M.Pd.\nJurusan Teknik Informatika dan Komputer - FT UNM (2026)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(2, 132, 199)

    # Real Interview Documentation Photo on Slide 1 (Cover)
    img_path = r'e:\SEM 3\DESAIN ANTAR MUKA PENGGUNA\SI-PRAS JTIK\foto_wawancara_w1.jpg'
    if os.path.exists(img_path):
        fr = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.8), Inches(4.9), Inches(3.2))
        fr.fill.solid()
        fr.fill.fore_color.rgb = C_WHITE
        fr.line.color.rgb = C_DARK
        fr.line.width = Pt(2)
        s1.shapes.add_picture(img_path, Inches(7.6), Inches(1.9), width=Inches(4.7), height=Inches(3.0))

        # Sticker badge below photo
        sb = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(5.15), Inches(4.9), Inches(0.6))
        sb.fill.solid(); sb.fill.fore_color.rgb = C_YELLOW; sb.line.color.rgb = C_DARK; sb.line.width = Pt(1.5)
        tf_sb = sb.text_frame; tf_sb.word_wrap = True
        p_sb = tf_sb.paragraphs[0]
        p_sb.text = "📸 Dokumentasi Wawancara W1 di Bengkel IT JTIK FT UNM"
        p_sb.font.name = FONT_FAMILY; p_sb.font.size = Pt(9.5); p_sb.font.bold = True; p_sb.font.color.rgb = C_DARK
        p_sb.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: BAB I - LATAR BELAKANG
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_base_slide(s2, "sisapras.jtik.unm.ac.id/latar-belakang", "BAB I")
    add_header(s2, "BAB I — PENDAHULUAN", "Latar Belakang & Kondisi Faktual JTIK", "Masalah pencatatan buku, sebaran lokasi, dan fenomena kursi ruang kelas")

    # Card 1
    add_card(s2, Inches(0.9), Inches(2.2), Inches(3.64), Inches(4.5), C_PINK)
    tb = s2.shapes.add_textbox(Inches(1.05), Inches(2.35), Inches(3.34), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Pencatatan Manual"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n[Kondisi Saat Ini]:\nPeminjaman barang masih banyak dicatat dengan menulis pada buku.\n\n[Dampak]:\nInformasi jumlah, ketersediaan, dan lokasi barang belum diketahui langsung.\n\n[Solusi UI]:\nLayanan terpadu digital satu pintu SI-SAPRAS JTIK."
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card 2
    add_card(s2, Inches(4.84), Inches(2.2), Inches(3.64), Inches(4.5), C_BLUE)
    tb = s2.shapes.add_textbox(Inches(4.99), Inches(2.35), Inches(3.34), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Sebaran Lokasi"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n[Kondisi Saat Ini]:\nBarang berada di beberapa tempat: Bengkel IT, laboratorium, dan ruang admin.\n\n[Dampak]:\nPengguna harus mencari atau menanyakan terlebih dahulu, atau mencari sendiri ke lokasi lain.\n\n[Solusi UI]:\nInformasi visual ketersediaan dan lokasi barang terpusat."
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card 3
    add_card(s2, Inches(8.78), Inches(2.2), Inches(3.64), Inches(4.5), C_GREEN)
    tb = s2.shapes.add_textbox(Inches(8.93), Inches(2.35), Inches(3.34), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Bangku Ruang Teori"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n[Kondisi Saat Ini]:\nJumlah kursi di ruang kelas sering tidak mencukupi kebutuhan perkuliahan.\n\n[Dampak]:\nPengguna mengambil/meminjam kursi dari kelas lain tanpa pencatatan.\n\n[Solusi UI]:\nMonitoring kuota kursi ruang kelas & alur lapor mutasi kursi."
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 3: BAB I - RUMUSAN MASALAH VS TUJUAN
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_base_slide(s3, "sisapras.jtik.unm.ac.id/fokus-penelitian", "PENYELARASAN")
    add_header(s3, "BAB I — FOKUS PENELITIAN", "Rumusan Masalah vs. Tujuan Penelitian", "Penyelarasan satu-ke-satu antara pertanyaan perancangan dan target capaian luaran")

    # Card Kiri
    add_card(s3, Inches(0.9), Inches(2.2), Inches(5.6), Inches(4.5), C_PURPLE)
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Rumusan Masalah (Pertanyaan Desain)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n1. Kebutuhan Informasi Pengguna:\n"
        "   Kebutuhan informasi apa yang dimiliki pengguna dalam mengetahui jumlah, ketersediaan, kondisi, dan lokasi sarpras sebelum meminjam?\n\n"
        "2. Alur Tugas dan Urutan Layar:\n"
        "   Bagaimana tugas mencari, meminjam, menggunakan, lapor pakai, dan mengembalikan diterjemahkan menjadi urutan layar SI-SAPRAS JTIK?\n\n"
        "3. Susunan Elemen Antarmuka:\n"
        "   Bagaimana kebutuhan informasi dan tugas diterjemahkan menjadi susunan antarmuka yang sesuai proses di JTIK?"
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card Kanan
    add_card(s3, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5), C_YELLOW)
    tb = s3.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Tujuan Penelitian (Target Capaian)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n1. Mengidentifikasi Kebutuhan Informasi:\n"
        "   Mengidentifikasi kebutuhan pengguna terkait jumlah, ketersediaan, kondisi, dan lokasi fisik barang.\n\n"
        "2. Menyusun Alur Tugas Pengguna:\n"
        "   Menyusun tugas mencari, meminjam, menggunakan, lapor pakai, dan mengembalikan menjadi alur layar SI-SAPRAS JTIK.\n\n"
        "3. Menghasilkan Rancangan Antarmuka:\n"
        "   Menghasilkan rancangan susunan informasi dan elemen antarmuka yang sesuai dengan proses peminjaman di JTIK."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 4: BAB I - BATASAN DAN RUANG LINGKUP
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_base_slide(s4, "sisapras.jtik.unm.ac.id/ruang-lingkup", "BATASAN")
    add_header(s4, "BAB I — RUANG LINGKUP", "Batasan dan Ruang Lingkup Perancangan", "Penetapan batas tegas tugas utama yang dirancang vs. fitur yang ditunda")

    # In-Scope (Expanded)
    add_card(s4, Inches(0.9), Inches(2.2), Inches(5.6), Inches(4.5), C_GREEN)
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Tugas Utama Dirancang (In-Scope)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• Mencari informasi barang, jumlah, dan ketersediaan stok.\n\n"
        "• Mengetahui lokasi barang di Bengkel IT, Lab, dan Ruang Admin.\n\n"
        "• Alur pengajuan peminjaman, pelaporan penggunaan, dan pengembalian.\n\n"
        "• Informasi inventaris yang berkaitan dengan proses tersebut."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Out-of-Scope (Expanded)
    add_card(s4, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5), C_PINK)
    tb = s4.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur yang Ditunda (Out-of-Scope)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• Dashboard perekaman aktivitas laboratorium otomatis.\n\n"
        "• Sensor penggunaan listrik di laboratorium.\n\n"
        "• Unit usaha untuk penyewaan barang komersial.\n\n"
        "• Pembuatan kode program, basis data, dan implementasi sistem langsung."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 5: BAB II - DESAIN UI VS IMK
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_base_slide(s5, "sisapras.jtik.unm.ac.id/teori-dasar", "BAB II")
    add_header(s5, "BAB II — TINJAUAN PUSTAKA", "Desain Antarmuka Pengguna vs. IMK", "Hubungan wujud visual di layar dengan payung keilmuan interaksi manusia-komputer")

    # Card UI
    add_card(s5, Inches(0.9), Inches(2.2), Inches(5.6), Inches(4.5), C_PINK)
    tb = s5.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Desain Antarmuka Pengguna (UI)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• Rujukan:\nAuliazmi et al. (2021); Lastiansah (2012)\n\n"
        "• Hakikat Konsep:\nSarana dialog dan interaksi antara manusia dengan komputer agar informasi disajikan secara mudah, menarik, dan komunikatif.\n\n"
        "• Komponen Layar:\nSemua hal yang terlihat di layar: tombol, formulir isian, status barang, dan navigasi menu.\n\n"
        "• Batasan Perancangan:\nDibatasi pada tampilan visual dan susunan informasi (tidak mencakup backend atau koding)."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card IMK
    add_card(s5, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5), C_BLUE)
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Interaksi Manusia–Komputer (IMK)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• Perbedaan dengan UI:\nIMK mencakup seluruh aspek hubungan antara manusia dan teknologi secara luas, termasuk kebiasaan dan cara berpikir pengguna.\n\n"
        "• Peran Desain UI:\nDesain antarmuka pengguna merupakan wujud nyata dari tampilan layar tempat interaksi manusia-komputer tersebut terjadi.\n\n"
        "• Penerapan pada SI-SAPRAS JTIK:\nMenyusun tampilan pencarian alat, pemetaan lokasi barang, dan alur peminjaman agar mudah dipahami pengguna."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 6: BAB II - HCD (ISO 9241-210)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_base_slide(s6, "sisapras.jtik.unm.ac.id/metode-hcd", "ISO 9241-210")
    add_header(s6, "BAB II — PENDEKATAN DESAIN", "Perancangan Berpusat pada Manusia (Human-Centered Design / HCD)", "Standar ISO 9241-210 (Wiryawan, 2011) dan 4 butir konteks penggunaan faktual di JTIK")

    # 4 Konteks Penggunaan Faktual JTIK
    contexts = [
        ("1. Pengguna (Users)", "Mahasiswa yang mencari & meminjam barang, Ketua Tingkat yang melapor peminjaman bangku, teknisi lab inventaris.", C_PINK),
        ("2. Tugas (Tasks)", "Mengecek ketersediaan alat, mengisi formulir peminjaman, melapor mutasi bangku kelas, konfirmasi pengembalian.", C_YELLOW),
        ("3. Peralatan (Equipment)", "Telepon pintar (smartphone) milik mahasiswa dan komputer meja (desktop PC) milik teknisi/pengelola lab.", C_BLUE),
        ("4. Lingkungan (Environment)", "Ruang kelas yang mendesak sebelum jam kuliah dimulai serta kondisi fisik laboratorium dan Bengkel IT.", C_GREEN)
    ]
    for idx, (ct_t, ct_d, ct_col) in enumerate(contexts):
        cx = Inches(0.9 + idx * 2.92)
        add_card(s6, cx, Inches(2.2), Inches(2.78), Inches(4.5), ct_col)
        tb = s6.shapes.add_textbox(cx + Inches(0.12), Inches(2.35), Inches(2.54), Inches(4.2))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ct_t
        p.font.name = FONT_FAMILY; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_DARK
        p = tf.add_paragraph()
        p.text = "\n" + ct_d
        p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 7: BAB II - HTA (DIX ET AL., 2004)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_base_slide(s7, "sisapras.jtik.unm.ac.id/struktur-tugas", "HTA & FLOW")
    add_header(s7, "BAB II — STRUKTUR TUGAS", "Hierarchical Task Analysis (HTA) & User Flow", "Dekomposisi tugas terstruktur dan alur perpindahan antarmuka (Dix et al., 2004)")

    # Card HTA
    add_card(s7, Inches(0.9), Inches(2.2), Inches(5.6), Inches(4.5), C_PURPLE)
    tb = s7.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Hierarchical Task Analysis (HTA)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• 4 Unsur Pokok HTA:\n"
        "  1. Goal (Indeks 0): Meminjam sarana praktikum laboratorium.\n"
        "  2. Sub-goals: 1. Cari ketersediaan, 2. Pilih lokasi, 3. Isi formulir.\n"
        "  3. Operations: Ketik nama barang pada pencarian, klik simpan.\n"
        "  4. Plan: Aturan urutan logis langkah yang dijalankan.\n\n"
        "• Stopping Rule (Aturan Berhenti):\n"
        "  Berhenti saat langkah sudah sederhana di layar tunggal. Ditandai garis tebal di bawah kotak (kedalaman maks 3 tingkat)."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card User Flow
    add_card(s7, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5), C_GREEN)
    tb = s7.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Alur Pengguna (User Flow)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• Peran pada SI-SAPRAS JTIK:\n"
        "  HTA digunakan untuk memetakan tugas utama pengguna.\n\n"
        "• Penerapan Alur Layar:\n"
        "  Urutan kegiatan peminjaman alat laboratorium di 3 ruangan serta pelaporan mutasi bangku kelas diatur secara jelas sebelum diwujudkan ke bentuk alur antarmuka (User Flow).\n\n"
        "• Sinergi Desain:\n"
        "  HTA menyusun pohon hierarki tugas, User Flow mewujudkan alur perpindahan visual antarlayar."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 8: BAB II - 5 KAJIAN PENELITIAN RELEVAN
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_base_slide(s8, "sisapras.jtik.unm.ac.id/penelitian-terdahulu", "5 JURNAL")
    add_header(s8, "BAB II — KAJIAN LITERATUR", "5 Kajian Antarmuka Relevan (State of the Art)", "Rujukan penelitian terdahulu yang mendasari perancangan SI-SAPRAS JTIK")

    papers = [
        ("1. Caylen Marli & Indah Lestari (2025)", "Facility Management System PCR. Segmentasi clustering manual, evaluasi SUS (77) & UEQ.", C_PINK),
        ("2. Cepeda & Saludes (2025)", "ICT Equipment Borrowing SPUP Filipina. Agile Scrum, ISO 25010 & TAM, rekomendasi responsivitas seluler.", C_BLUE),
        ("3. Setya Hadi Nugroho & Waris Marsisno (2025)", "Peminjaman Barang & Ruang Polstat STIS. SDLC Prototyping, pengujian Black Box & survei PSSUQ.", C_GREEN),
        ("4. Waliadi Gunawan, Abdul Ghofur, dkk. (2025)", "Monitoring Alat Lab Akbid Bina Husada. Waterfall, PHP & MySQL, kendala padam listrik/offline.", C_YELLOW),
        ("5. Dimas Jayadi & Ucuk Darusalam (2022)", "Sistem Peminjaman Alat Lab Android. FAST Framework & Firebase, untuk asisten lab Computer Vision.", C_PURPLE)
    ]
    for idx, (p_t, p_d, p_bg) in enumerate(papers):
        cy = Inches(2.2 + idx * 0.9)
        add_card(s8, Inches(0.9), cy, Inches(11.5), Inches(0.8), p_bg)
        tb = s8.shapes.add_textbox(Inches(1.1), cy + Inches(0.08), Inches(11.1), Inches(0.65))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{p_t}  ➔  {p_d}"
        p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 9: BAB III - METODE PENELITIAN & NARASUMBER
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_base_slide(s9, "sisapras.jtik.unm.ac.id/metodologi", "BAB III")
    add_header(s9, "BAB III — METODE PENELITIAN", "Pendekatan Perancangan, Teknik, & Narasumber", "Pengumpulan data empiris melalui wawancara mendalam di Bengkel IT JTIK FT UNM")

    # Card 1
    add_card(s9, Inches(0.9), Inches(2.2), Inches(3.64), Inches(4.5), C_PINK)
    tb = s9.shapes.add_textbox(Inches(1.05), Inches(2.35), Inches(3.34), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Pendekatan Desain"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n• Metode:\nHuman-Centered Design (HCD).\n\n• Alasan Pemilihan:\nBerangkat dari masalah nyata di JTIK agar rancangan antarmuka didasarkan pada kebutuhan empiris pengguna.\n\n• Fokus Proposal UTS:\nPengumpulan data wawancara, analisis kebutuhan, penyusunan HTA & User Flow, serta perancangan antarmuka."
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Card 2
    add_card(s9, Inches(4.84), Inches(2.2), Inches(3.64), Inches(4.5), C_YELLOW)
    tb = s9.shapes.add_textbox(Inches(4.99), Inches(2.35), Inches(3.34), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Teknik Pengumpulan Data"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n• Wawancara Mendalam:\nWawancara langsung untuk menggali prosedur operasional dan hambatan nyata dalam pengelolaan sarpras JTIK.\n\n• Fokus Penggalian:\nKendala pencatatan buku konvensional, sebaran alat di 3 ruangan, fenomena bangku kelas, dan verifikasi bebas sarpras.\n\n• Tujuan Data:\nMenjadi landasan penyusunan HTA, User Flow, serta arsitektur solusi sistem informasi."
    p.font.name = FONT_FAMILY; p.font.size = Pt(10.5); p.font.color.rgb = C_DARK

    # Card 3
    add_card(s9, Inches(8.78), Inches(2.2), Inches(3.64), Inches(4.5), C_GREEN)
    tb = s9.shapes.add_textbox(Inches(8.93), Inches(2.35), Inches(3.34), Inches(2.3))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Narasumber (W1)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "\n• Kode: W1 (Kepala Lab)\n• Metode: Wawancara Mendalam\n• Lokasi: Bengkel IT JTIK FT UNM\n• Tanggal: 6 Oktober 2026\n• Etika: Persetujuan & Penyamaran W1"
    p.font.name = FONT_FAMILY; p.font.size = Pt(10.5); p.font.color.rgb = C_DARK

    if os.path.exists(img_path):
        s9.shapes.add_picture(img_path, Inches(8.93), Inches(4.75), width=Inches(3.34), height=Inches(1.8))

    # ==========================================
    # SLIDE 10: BAB III - 5 PERTANYAAN WAWANCARA (3.4)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_base_slide(s10, "sisapras.jtik.unm.ac.id/instrumen", "INSTRUMEN")
    add_header(s10, "BAB III — INSTRUMEN PENELITIAN", "3.4 Instrumen Pengumpulan Data Wawancara", "Lima butir pertanyaan wawancara mendalam dengan Kepala Laboratorium (W1)")

    add_card(s10, Inches(0.9), Inches(2.2), Inches(11.5), Inches(4.5), C_WHITE)
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(11.1), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Lima Pertanyaan Pokok Pedoman Wawancara:"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n1. Pendataan Inventaris:\n"
        "   Bagaimana proses pengecekan dan pendataan alat serta bahan yang menjadi inventaris laboratorium saat ini?\n\n"
        "2. Lokasi Penyimpanan Barang:\n"
        "   Bagaimana cara mengetahui lokasi penyimpanan setiap barang atau alat yang ada di laboratorium?\n\n"
        "3. Bukti Penggunaan Alat:\n"
        "   Bagaimana proses pencatatan atau pemberian bukti penggunaan alat dan bahan setelah digunakan oleh pengguna?\n\n"
        "4. Pelaporan Penggunaan:\n"
        "   Bagaimana proses pelaporan penggunaan alat dilakukan setelah peminjaman selesai?\n\n"
        "5. Bukti Bebas Peminjaman:\n"
        "   Bagaimana proses penerbitan atau pencetakan bukti bebas peminjaman alat saat ini?"
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 11: BAB III - ANALISIS DATA & PENUTUP
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_base_slide(s11, "sisapras.jtik.unm.ac.id/closing", "PENUTUP")
    add_header(s11, "BAB III — PENUTUP", "Sesi Tanya Jawab", "")

    # Kiri: Analisis Data
    add_card(s11, Inches(0.9), Inches(2.2), Inches(5.6), Inches(4.5), C_BLUE)
    tb = s11.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Teknik Analisis Data (Deskriptif)"
    p.font.name = FONT_FAMILY; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\n• 1. Telaah Catatan Wawancara:\n"
        "  Informasi dikelompokkan berdasarkan permasalahan, aktivitas, dan proses pengelolaan serta peminjaman sarpras.\n\n"
        "• 2. Perumusan Kebutuhan Pengguna:\n"
        "  Temuan tidak langsung dituliskan sebagai fitur, tetapi dirumuskan berdasarkan apa yang perlu diketahui/dilakukan pengguna beserta alasannya.\n\n"
        "• 3. Skala Prioritas:\n"
        "  Diprioritaskan berdasarkan frekuensi kemunculan dalam wawancara dan dampak yang terjadi jika tidak terpenuhi."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(11); p.font.color.rgb = C_DARK

    # Kanan: Penutup / Q&A
    add_card(s11, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5), C_YELLOW)
    tb = s11.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.2), Inches(4.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Terima Kasih!"
    p.font.name = FONT_FAMILY; p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = (
        "\nSesi Diskusi, Masukan, dan Tanya Jawab\n\n"
        "Kelompok 2 (HCI) — Program Studi Teknik Komputer\n"
        "Jurusan Teknik Informatika dan Komputer (JTIK)\n"
        "Universitas Negeri Makassar (2026)\n\n"
        "Dosen Pengampu: Ayu Lestari, S.Pd., M.Pd."
    )
    p.font.name = FONT_FAMILY; p.font.size = Pt(12); p.font.color.rgb = C_DARK

    output_path = r'e:\SEM 3\DESAIN ANTAR MUKA PENGGUNA\SI-PRAS JTIK\Presentasi_UTS_SI-SAPRAS_JTIK_FotoWawancara.pptx'
    try:
        prs.save(r'e:\SEM 3\DESAIN ANTAR MUKA PENGGUNA\SI-PRAS JTIK\Presentasi_UTS_SI-SAPRAS_JTIK_Final.pptx')
        print("Updated Presentasi_UTS_SI-SAPRAS_JTIK_Final.pptx")
    except Exception as e:
        print(f"File locked by PowerPoint: {e}")
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == '__main__':
    build_final_presentation()
